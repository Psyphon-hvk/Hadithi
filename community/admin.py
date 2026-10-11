
from io import BytesIO
from xml.sax.saxutils import escape

from django.contrib import admin
from django.http import HttpResponse
from django.utils import timezone

from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import getSampleStyleSheet, ParagraphStyle
from reportlab.lib.units import inch
from reportlab.lib import colors
from reportlab.platypus import (
    SimpleDocTemplate,
    Paragraph,
    Spacer,
    PageBreak,
    HRFlowable,
)

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter

from .models import (
    Assessment,
    AssessmentQuestion,
    AssessmentResult,
    Comment,
    Story,
)


class AssessmentQuestionInline(admin.TabularInline):
    model = AssessmentQuestion
    extra = 1


# =========================================================
# PDF EXPORT
# =========================================================

@admin.action(description="Download selected stories + conversations (PDF)")
def export_stories_as_pdf(modeladmin, request, queryset):
    """
    Export selected stories and their full comment threads
    as a single PDF report.
    """
    response = HttpResponse(content_type="application/pdf")

    filename = (
        f"hadithi_stories_export_"
        f"{timezone.now().strftime('%Y%m%d_%H%M%S')}.pdf"
    )

    response["Content-Disposition"] = (
        f'attachment; filename="{filename}"'
    )

    doc = SimpleDocTemplate(
        response,
        pagesize=letter,
        topMargin=0.75 * inch,
        bottomMargin=0.75 * inch,
        leftMargin=0.75 * inch,
        rightMargin=0.75 * inch,
    )

    styles = getSampleStyleSheet()

    title_style = styles["Title"]

    heading_style = ParagraphStyle(
        "StoryHeading",
        parent=styles["Heading2"],
        spaceBefore=6,
        spaceAfter=4,
    )

    meta_style = ParagraphStyle(
        "Meta",
        parent=styles["Normal"],
        textColor=colors.grey,
        fontSize=9,
        spaceAfter=8,
    )

    body_style = ParagraphStyle(
        "Body",
        parent=styles["Normal"],
        spaceAfter=10,
        leading=14,
    )

    comment_heading_style = ParagraphStyle(
        "CommentHeading",
        parent=styles["Heading3"],
        spaceBefore=10,
        spaceAfter=6,
    )

    comment_style = ParagraphStyle(
        "Comment",
        parent=styles["Normal"],
        leftIndent=14,
        spaceAfter=6,
        leading=13,
    )

    comment_meta_style = ParagraphStyle(
        "CommentMeta",
        parent=styles["Normal"],
        leftIndent=14,
        textColor=colors.grey,
        fontSize=8,
        spaceAfter=2,
    )

    elements = []

    stories = list(
        queryset.order_by("-created_at").prefetch_related(
            "comments__author",
            "likes",
        )
    )

    elements.append(
        Paragraph(
            "Hadithi - Stories &amp; Conversations Export",
            title_style,
        )
    )

    elements.append(
        Paragraph(
            f"Generated {timezone.localtime(timezone.now()).strftime('%d %b %Y, %H:%M')} "
            f"· {len(stories)} storie(s) included",
            meta_style,
        )
    )

    elements.append(Spacer(1, 12))

    for index, story in enumerate(stories):
        if index > 0:
            elements.append(PageBreak())

        elements.append(
            Paragraph(escape(str(story.title)), heading_style)
        )

        status = (
            "Approved"
            if story.is_approved
            else "Pending moderation"
        )

        posted_date = timezone.localtime(
            story.created_at
        ).strftime("%d %b %Y, %H:%M")

        elements.append(
            Paragraph(
                f"By: {escape(str(story.display_author))} "
                f"&nbsp;|&nbsp; Posted: {posted_date} "
                f"&nbsp;|&nbsp; Status: {status} "
                f"&nbsp;|&nbsp; Likes: {story.likes.count()}",
                meta_style,
            )
        )

        elements.append(
            HRFlowable(
                width="100%",
                color=colors.lightgrey,
                thickness=0.5,
            )
        )

        elements.append(Spacer(1, 6))

        # Story body, preserving paragraph breaks.
        for para in story.body.splitlines():
            if para.strip():
                elements.append(
                    Paragraph(
                        escape(para.strip()),
                        body_style,
                    )
                )

        comments = list(
            story.comments.select_related("author").order_by(
                "created_at"
            )
        )

        comment_count = len(comments)

        elements.append(
            Paragraph(
                f"Conversation ({comment_count} "
                f"comment{'s' if comment_count != 1 else ''})",
                comment_heading_style,
            )
        )

        if not comments:
            elements.append(
                Paragraph(
                    "No comments yet.",
                    comment_meta_style,
                )
            )
        else:
            for comment in comments:
                comment_date = timezone.localtime(
                    comment.created_at
                ).strftime("%d %b %Y, %H:%M")

                elements.append(
                    Paragraph(
                        f"{escape(str(comment.author))} "
                        f"· {comment_date}",
                        comment_meta_style,
                    )
                )

                for para in comment.body.splitlines():
                    if para.strip():
                        elements.append(
                            Paragraph(
                                escape(para.strip()),
                                comment_style,
                            )
                        )

    doc.build(elements)

    return response


# =========================================================
# EXCEL EXPORT
# =========================================================

@admin.action(description="Download selected stories + conversations (Excel)")
def export_stories_as_excel(modeladmin, request, queryset):
    """
    Export selected stories and their comment threads
    into an Excel workbook with Stories and Comments sheets.
    """
    workbook = Workbook()

    # -----------------------------------------------------
    # STORIES WORKSHEET
    # -----------------------------------------------------

    stories_sheet = workbook.active
    stories_sheet.title = "Stories"

    stories_sheet.append([
        "Story ID",
        "Title",
        "Author",
        "Anonymous",
        "Approved",
        "Date Posted",
        "Likes",
        "Story Content",
    ])

    # -----------------------------------------------------
    # COMMENTS WORKSHEET
    # -----------------------------------------------------

    comments_sheet = workbook.create_sheet("Comments")

    comments_sheet.append([
        "Story ID",
        "Story Title",
        "Comment Author",
        "Comment Date",
        "Comment Content",
    ])

    # -----------------------------------------------------
    # STYLE WORKSHEETS
    # -----------------------------------------------------

    header_fill = PatternFill(
        fill_type="solid",
        fgColor="263238",
    )

    header_font = Font(
        bold=True,
        color="FFFFFF",
        size=11,
    )

    for sheet in (stories_sheet, comments_sheet):
        sheet.freeze_panes = "A2"
        sheet.sheet_view.showGridLines = False

        for cell in sheet[1]:
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(
                horizontal="center",
                vertical="center",
                wrap_text=True,
            )

        sheet.row_dimensions[1].height = 30

    # -----------------------------------------------------
    # EXPORT SELECTED STORIES AND COMMENTS
    # -----------------------------------------------------

    stories = queryset.order_by(
        "-created_at"
    ).prefetch_related(
        "comments__author",
        "likes",
    )

    for story in stories:
        stories_sheet.append([
            story.pk,
            str(story.title),
            str(story.display_author),
            "Yes" if story.is_anonymous else "No",
            "Approved" if story.is_approved else "Pending moderation",
            timezone.localtime(
                story.created_at
            ).strftime("%d %b %Y, %H:%M"),
            story.likes.count(),
            str(story.body),
        ])

        story_row = stories_sheet.max_row

        stories_sheet.cell(
            row=story_row,
            column=8,
        ).alignment = Alignment(
            vertical="top",
            wrap_text=True,
        )

        # Keep each comment linked to its parent story by ID.
        comments = story.comments.select_related(
            "author"
        ).order_by("created_at")

        for comment in comments:
            comments_sheet.append([
                story.pk,
                str(story.title),
                str(comment.author),
                timezone.localtime(
                    comment.created_at
                ).strftime("%d %b %Y, %H:%M"),
                str(comment.body),
            ])

            comment_row = comments_sheet.max_row

            for column in (2, 3, 5):
                comments_sheet.cell(
                    row=comment_row,
                    column=column,
                ).alignment = Alignment(
                    vertical="top",
                    wrap_text=True,
                )

    # -----------------------------------------------------
    # COLUMN WIDTHS
    # -----------------------------------------------------

    story_widths = {
        1: 12,
        2: 35,
        3: 25,
        4: 14,
        5: 24,
        6: 22,
        7: 10,
        8: 70,
    }

    comment_widths = {
        1: 12,
        2: 35,
        3: 25,
        4: 22,
        5: 70,
    }

    for column, width in story_widths.items():
        stories_sheet.column_dimensions[
            get_column_letter(column)
        ].width = width

    for column, width in comment_widths.items():
        comments_sheet.column_dimensions[
            get_column_letter(column)
        ].width = width

    # Add filters to the complete datasets.
    for sheet in (stories_sheet, comments_sheet):
        if sheet.max_row > 1:
            sheet.auto_filter.ref = sheet.dimensions
        else:
            sheet.auto_filter.ref = (
                f"A1:{get_column_letter(sheet.max_column)}1"
            )

    # -----------------------------------------------------
    # GENERATE DOWNLOADABLE WORKBOOK
    # -----------------------------------------------------

    output = BytesIO()
    workbook.save(output)
    output.seek(0)

    filename = (
        f"hadithi_stories_export_"
        f"{timezone.now().strftime('%Y%m%d_%H%M%S')}.xlsx"
    )

    response = HttpResponse(
        output.getvalue(),
        content_type=(
            "application/vnd.openxmlformats-officedocument."
            "spreadsheetml.sheet"
        ),
    )

    response["Content-Disposition"] = (
        f'attachment; filename="{filename}"'
    )

    return response


# =========================================================
# DJANGO ADMIN REGISTRATION
# =========================================================

@admin.register(Assessment)
class AssessmentAdmin(admin.ModelAdmin):
    list_display = (
        "name",
        "assessment_type",
        "is_active",
    )

    inlines = [AssessmentQuestionInline]


@admin.register(Story)
class StoryAdmin(admin.ModelAdmin):
    list_display = (
        "title",
        "author",
        "is_anonymous",
        "is_approved",
        "created_at",
    )

    list_filter = (
        "is_anonymous",
        "is_approved",
    )

    prepopulated_fields = {
        "slug": ("title",),
    }

    actions = [
        export_stories_as_pdf,
        export_stories_as_excel,
    ]


@admin.register(Comment)
class CommentAdmin(admin.ModelAdmin):
    list_display = (
        "story",
        "author",
        "created_at",
    )


@admin.register(AssessmentResult)
class AssessmentResultAdmin(admin.ModelAdmin):
    list_display = (
        "user",
        "assessment",
        "score",
        "taken_at",
    )

    list_filter = ("assessment",)
